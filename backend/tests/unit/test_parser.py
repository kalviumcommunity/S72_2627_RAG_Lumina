from __future__ import annotations

from pathlib import Path

import pytest

from app.services.ingestion.parser import (
    ParseError,
    blocks_from_plain_pages,
    parse_file,
    parse_html,
    parse_markdown,
)


def test_markdown_blocks_and_offsets() -> None:
    text = "# Title\n\nIntro line one\nline two\n\n- item one\n- item two\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
    doc = parse_markdown(text)
    kinds = [b.type for b in doc.blocks]
    assert kinds == ["heading", "paragraph", "list_item", "list_item", "table"]
    for block in doc.blocks:
        source = doc.full_text[block.char_start : block.char_end]
        # Headings store their title without the "#" markup; everything else is the verbatim slice.
        assert source.endswith(block.text) if block.type == "heading" else source == block.text
    assert doc.blocks[-1].rows == [["A", "B"], ["1", "2"]]


def test_plain_text_headings_are_detected_but_wrapped_sentences_are_not() -> None:
    page = (
        "1 BACKGROUND\nAn audit found a problem.\n\n2 AMENDMENT\n"
        "Section 5.3 of P-ED-01 (Sepsis Recognition and One-Hour Bundle, version 2) is\n"
        "replaced with effect from 15 September 2026.\n(a) Give antibiotics within 1 hour.\n"
    )
    blocks, full = blocks_from_plain_pages([page])
    headings = [b.text for b in blocks if b.type == "heading"]
    assert headings == ["1 BACKGROUND", "2 AMENDMENT"]
    assert any(b.type == "list_item" and b.text.startswith("(a)") for b in blocks)
    for block in blocks:
        assert full[block.char_start : block.char_end].strip() == block.text.strip()


def test_html_parsing() -> None:
    html = "<html><body><nav>skip</nav><h2>3 Scope</h2><p>Adults only.</p><ul><li>ED</li></ul>"
    html += "<table><tr><th>Drug</th><th>Class</th></tr><tr><td>Meropenem</td><td>Restricted</td></tr></table></body></html>"
    doc = parse_html(html)
    assert [b.type for b in doc.blocks] == ["heading", "paragraph", "list_item", "table"]
    assert "skip" not in doc.full_text
    assert doc.blocks[-1].rows == [["Drug", "Class"], ["Meropenem", "Restricted"]]


def test_docx_parsing(tmp_path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("4.2 Nomogram", level=2)
    document.add_paragraph("Adjust according to the table.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "aPTT", "Action"
    table.cell(1, 0).text, table.cell(1, 1).text = "Above 100", "Hold"
    path = tmp_path / "doc.docx"
    document.save(path)
    doc = parse_file(path)
    assert [b.type for b in doc.blocks] == ["heading", "paragraph", "table"]
    assert doc.blocks[0].text == "4.2 Nomogram"
    assert doc.blocks[2].rows == [["aPTT", "Action"], ["Above 100", "Hold"]]


def test_unsupported_extension_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "evil.exe"
    path.write_bytes(b"MZ")
    with pytest.raises(ParseError):
        parse_file(path)
